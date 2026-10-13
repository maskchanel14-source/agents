#!/usr/bin/env python3
"""Audit the built PPTX against spec.json (geometry, containment, accessibility).

Implements the pptx-quality-gates checklist: reopen the package, compare actual
bounds with the specification, check collisions, safe margins, type scale, text
capacity, table headers, and reading order. Writes audits/ JSON reports.

Usage:  python3 audit_deck.py [spec.json] [deck.pptx]
Exit code is non-zero when any check fails.
"""
from __future__ import annotations

import json
import math
import sys
from datetime import datetime, timezone
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

HERE = Path(__file__).resolve().parent
EMU_PER_INCH = 914400
TOLERANCE_IN = 0.02
TEXT_ROLES = {"title", "kicker", "body", "subtitle", "meta", "presenter", "callout",
              "takeaway", "panel_text", "metric", "footer", "table"}
SKIP_COLLISION_ROLES = {"bar", "background", "rule"}


def inches(value: int) -> float:
    return round(value / EMU_PER_INCH, 3)


def overlap(a, b, gap=0.0):
    return (a["x"] < b["x"] + b["width"] - gap and b["x"] < a["x"] + a["width"] - gap
            and a["y"] < b["y"] + b["height"] - gap and b["y"] < a["y"] + a["height"] - gap)


def contains(outer, inner, pad=0.0):
    return (inner["x"] >= outer["x"] - pad and inner["y"] >= outer["y"] - pad
            and inner["x"] + inner["width"] <= outer["x"] + outer["width"] + pad
            and inner["y"] + inner["height"] <= outer["y"] + outer["height"] + pad)


def estimate_lines(text: str, width_in: float, font_pt: float) -> int:
    usable = max(0.1, width_in - 0.1)
    chars_per_line = max(8, int(usable * 96 / (0.53 * font_pt)))
    lines = 0
    for chunk in text.split("\n"):
        lines += max(1, math.ceil(len(chunk) / chars_per_line))
    return lines


def main() -> None:
    spec_path = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "spec.json"
    deck_path = Path(sys.argv[2]) if len(sys.argv) > 2 else HERE / "Presentasi-Jurnal-MIH-Lopes-2021.pptx"
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    prs = Presentation(str(deck_path))
    policy = spec["summary"]["layout_policy"]
    slide_w = spec["slides"][0]["layout_tree"]["slide_size"]["width"]
    slide_h = spec["slides"][0]["layout_tree"]["slide_size"]["height"]

    findings: list[dict] = []
    line_exceptions: set[tuple[str, str]] = set()
    geometry: list[dict] = []
    accessibility: list[dict] = []

    def flag(kind, severity, slide_id, object_id, message, fix):
        findings.append({"kind": kind, "severity": severity, "slide_id": slide_id,
                         "object_id": object_id, "message": message, "fix": fix})

    if len(prs.slides) != len(spec["slides"]):
        flag("SLIDE_COUNT", "error", "-", "-",
             f"deck has {len(prs.slides)} slides, spec has {len(spec['slides'])}", "rebuild the deck")

    for slide_spec, slide in zip(spec["slides"], prs.slides):
        sid = slide_spec["id"]
        declared = slide_spec["layout_tree"]["objects"]
        actual: dict[str, dict] = {}
        shape_index = {}
        for index, shape in enumerate(slide.shapes):
            actual[f"__shape_{index}"] = {
                "x": inches(shape.left), "y": inches(shape.top),
                "width": inches(shape.width), "height": inches(shape.height),
                "name": shape.shape_type.name if shape.shape_type is not None else "unknown",
            }
            shape_index[index] = shape

        if getattr(slide, "element", None) is not None and slide.slide_id is None:  # pragma: no cover
            flag("SLIDE_ID", "warning", sid, "-", "slide has no id", "verify deck integrity")

        # 1) geometry comparison, object by object, in declared z-order
        ordered = sorted(declared.values(), key=lambda o: (o["z_index"], o["id"]))
        if len(ordered) != len(slide.shapes):
            flag("OBJECT_COUNT", "error", sid, "-",
                 f"{len(slide.shapes)} shapes in pptx vs {len(ordered)} objects in spec",
                 "rebuild; do not hand-edit the deck")
        for spec_obj, shape in zip(ordered, slide.shapes):
            b = spec_obj["bbox"]
            got = {"x": inches(shape.left), "y": inches(shape.top),
                   "width": inches(shape.width), "height": inches(shape.height)}
            keys = ("x", "y", "width") if spec_obj["kind"] == "line" else ("x", "y", "width", "height")
            deviation = max(abs(got[k] - b[k]) for k in keys)
            geometry.append({"slide_id": sid, "object_id": spec_obj["id"], "declared": b,
                             "actual": got, "deviation_in": round(deviation, 4),
                             "compared_fields": list(keys)})
            if deviation > TOLERANCE_IN:
                flag("GEOMETRY_DRIFT", "error", sid, spec_obj["id"],
                     f"declared {b} vs actual {got}", "fix spec or builder, rebuild")
            if spec_obj["kind"] == "line":
                # documented exception: a straight 1.5 pt connector normalises to zero
                # height in OOXML, so the declared 0.028 in is a nominal stroke box.
                line_exceptions.add((sid, spec_obj["id"]))

        # 2) positive geometry, safe margins, footer rail
        for o in declared.values():
            b = o["bbox"]
            if b["width"] <= 0 or b["height"] <= 0:
                flag("NON_POSITIVE_GEOMETRY", "error", sid, o["id"], "zero/negative size", "resize")
            if o["classification"] == "layout_design":
                continue
            if b["x"] < policy["safe_margin"] - 0.01 or b["x"] + b["width"] > slide_w - policy["safe_margin"] + 0.01:
                flag("UNSAFE_X", "error", sid, o["id"], f"x range {b['x']}..{b['x']+b['width']}",
                     "move inside the safe margin")
            if o["role"] == "footer":
                continue
            if b["y"] + b["height"] > policy["content_bottom"] + 0.01:
                flag("UNSAFE_BOTTOM", "error", sid, o["id"],
                     f"bottom {round(b['y']+b['height'],3)} exceeds content area",
                     "shorten or move content above the footer rail")

        # 3) collisions between text-bearing objects (card/panel containment allowed)
        texty = [o for o in declared.values() if o["role"] not in SKIP_COLLISION_ROLES
                 and o["classification"] == "content"]
        for i, a in enumerate(texty):
            for c in texty[i + 1:]:
                if not overlap(a["bbox"], c["bbox"], gap=0.02):
                    continue
                if contains(a["bbox"], c["bbox"]) or contains(c["bbox"], a["bbox"]):
                    continue  # one sits inside the other's box on purpose
                flag("CONTENT_COLLISION", "error", sid, f"{a['id']}|{c['id']}",
                     "overlapping content bboxes", "move or resize one object")

        # 4) text capacity + containment, type scale
        for o in declared.values():
            if o["kind"] not in ("text", "table"):
                continue
            b, style = o["bbox"], o["style"]
            if o["kind"] == "text":
                size = float(style.get("font_size", 12))
                if o["role"] != "footer" and o["classification"] == "content" and size < 9:
                    flag("TYPE_SCALE", "error", sid, o["id"], f"{size} pt below the 9 pt floor", "raise size")
                lines = estimate_lines(o["content"]["text"], b["width"], size)
                needed = lines * (size * 1.22 / 72) + 0.04
                if needed > b["height"] + 0.06:
                    flag("TEXT_CAPACITY", "error", sid, o["id"],
                         f"needs ~{round(needed,2)}in for {lines} lines, box is {b['height']}in",
                         "shorten copy, the special case: grow the box via the spec")
                # containment in a card/panel when one encloses the text box
                parents = [p for p in declared.values()
                           if p["role"] in ("card", "panel") and contains(p["bbox"], b)]
                if parents:
                    parent = min(parents, key=lambda p: p["bbox"]["width"] * p["bbox"]["height"])
                    pad = 0.12
                    inner = {"x": parent["bbox"]["x"] + pad, "y": parent["bbox"]["y"] + pad,
                             "width": parent["bbox"]["width"] - 2 * pad,
                             "height": parent["bbox"]["height"] - 2 * pad}
                    if not contains(inner, b):
                        flag("CONTAINMENT", "error", sid, o["id"],
                             f"text box exceeds inner padding of {parent['id']}",
                             "shrink the text bbox")
            else:
                widths = sum(style["col_widths"])
                if abs(widths - b["width"]) > 0.05:
                    flag("TABLE_WIDTH", "error", sid, o["id"],
                         f"columns sum to {round(widths,3)}in, bbox is {b['width']}in",
                         "recompute column widths")
                if len(o["content"]["header"]) != len(style["col_widths"]):
                    flag("TABLE_HEADER", "error", sid, o["id"], "header/col mismatch",
                         "align header with columns")

        # 5) accessibility: reading order, titles, table headers
        ids = set(declared)
        for ro_id in slide_spec["accessibility"]["reading_order"]:
            if ro_id not in ids:
                flag("READING_ORDER", "error", sid, ro_id, "reading-order id not found in the tree",
                     "fix the reading order list")
        if "title" not in declared:
            flag("SLIDE_TITLE", "error", sid, "-", "slide has no title object", "add a title object")
        accessibility.append({"slide_id": sid, "title": slide_spec["accessibility"]["title"],
                              "reading_order": slide_spec["accessibility"]["reading_order"],
                              "objects": len(declared)})

    errors = [f for f in findings if f["severity"] == "error"]
    exceptions = [{"kind": "LINE_CONNECTOR_NORMALISATION", "slide_id": sid, "object_id": oid,
                   "reason": "straight connector renders at zero OOXML height; declared 0.028 in is a "
                             "nominal stroke box, stroke width 1.5 pt",
                   "owner": "deck author", "review_date": "next deck revision"}
                  for sid, oid in sorted(line_exceptions)]
    report = {"audited_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
              "deck": deck_path.name, "spec": spec_path.name,
              "slides": len(prs.slides), "objects_checked": len(geometry),
              "accepted_exceptions": exceptions,
              "checks": ["geometry-drift", "positive-geometry", "safe-margins", "footer-rail",
                         "content-collisions", "text-capacity", "containment", "type-scale",
                         "table-integrity", "reading-order", "slide-titles"],
              "errors": len(errors), "findings": findings}
    (HERE / "audits").mkdir(exist_ok=True)
    (HERE / "audits" / "geometry-audit.json").write_text(
        json.dumps({**report, "samples": geometry[:40]}, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8")
    (HERE / "audits" / "accessibility-audit.json").write_text(
        json.dumps({"audited_at": report["audited_at"], "language": "id-ID",
                    "slides": accessibility}, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8")
    (HERE / "audits" / "quality-gates.json").write_text(
        json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    for f in errors:
        print(f"ERROR   {f['kind']:<20} {f['slide_id']:<18} {f['object_id']:<22} {f['message']}")
    print(f"audit done: {report['objects_checked']} objects, {len(errors)} error(s)")
    raise SystemExit(1 if errors else 0)


if __name__ == "__main__":
    main()
